from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from src.reporting.periods import bounds, last_month, TZ, utc
from src.reporting.store import ReportStore
from src.reporting.schema import validate_findings
from src.reporting.pipeline import service_image, chunks, MonthlyService
from src.reporting.ews import EWSClient, NS
from src.reporting.routes import router


def task(**kwargs):
    return dict(
        id="t1",
        section="4.2",
        title="Планировка Suite",
        result="Направлена на согласование.",
        status="На согласовании",
        previous_id=None,
        evidence_ids=[],
        attachment_ids=[],
        include_in_report=True,
        **kwargs,
    )


def findings(tasks=None):
    return {"tasks": tasks or [], "risks": [], "warnings": []}


def test_periods_and_timezone():
    now = datetime(2026, 9, 1, tzinfo=TZ)
    assert last_month(now) == "2026-08"
    start, end = bounds("2026-08", now)
    assert utc(start) == "2026-07-31T21:00:00Z"
    assert utc(end) == "2026-08-31T21:00:00Z"
    with pytest.raises(ValueError):
        bounds("2026-09", now)
    assert bounds("2024-02", now)[1].day == 1


def test_approved_history_immutable_and_previous_skips_drafts(tmp_path):
    store = ReportStore(tmp_path / "history.sqlite")
    july = store.create("2026-07", status="draft")
    july = store.update(july["id"], status="approved")
    store.create("2026-08", status="draft")
    assert store.previous("2026-09")["id"] == july["id"]
    assert store.previous("2026-07") is None
    with pytest.raises(ValueError):
        store.update(july["id"], tasks=[])


def test_revision_conflict_and_restart(tmp_path):
    store = ReportStore(tmp_path / "history.sqlite")
    row = store.create("2026-08")
    store.update(row["id"], status="processing")
    with pytest.raises(ValueError):
        store.update(row["id"], expected_revision=1, status="draft")
    recovered = ReportStore(store.path).get(row["id"])
    assert recovered["status"] == "failed"
    assert "перезапуск" in recovered["error"]


def test_no_news_cannot_close_task():
    old = task()
    new = dict(
        old, id="temporary", previous_id="t1", status="Завершено", result="Согласовано"
    )
    result = validate_findings(findings([new]), [], findings([old]))
    assert result["tasks"][0]["status"] == old["status"]
    assert result["tasks"][0]["result"] == old["result"]
    assert result["tasks"][0]["include_in_report"] is False
    assert result["warnings"]


def test_closed_task_retained_without_repetition_and_can_reopen():
    old = dict(task(), status="Завершено")
    result = validate_findings(findings(), [], findings([old]))
    assert len(result["tasks"]) == 1
    assert not result["tasks"][0]["include_in_report"]
    assert not result["warnings"]
    new = dict(
        task(), id="newid", previous_id="t1", status="В работе", evidence_ids=["m1"]
    )
    result = validate_findings(
        findings([new]), [{"id": "m1", "attachments": []}], findings([old])
    )
    assert result["tasks"][0]["id"] == "t1"
    assert result["tasks"][0]["status"] == "В работе"


def test_invented_evidence_and_unrelated_attachments_rejected():
    new = dict(task(), evidence_ids=["made-up"])
    with pytest.raises(ValueError):
        validate_findings(findings([new]), [])
    new.update(evidence_ids=["m1"], attachment_ids=["from-other-mail"])
    with pytest.raises(ValueError):
        validate_findings(
            findings([new]), [{"id": "m1", "attachments": [{"id": "a1"}]}]
        )


def test_images_are_not_discarded_by_extension_alone():
    assert service_image({"name": "logo.png", "inline": False, "size": 90000})
    assert service_image({"name": "image001.jpg", "inline": True, "size": 9000})
    assert not service_image(
        {"name": "Site_photo.jpg", "inline": False, "size": 900000}
    )
    assert not service_image({"name": "drawing.png", "inline": True, "size": 900000})
    assert not service_image(
        {"name": "specification.pdf", "inline": True, "size": 10000}
    )


def test_long_mail_all_text_reaches_analysis():
    body = "A" * 30000 + "important final decision"
    batches = list(chunks([{"id": "m1", "body": body}], limit=16000))
    assert "".join(x["body"] for b in batches for x in b) == body
    assert all(x["id"] == "m1" for b in batches for x in b)


def test_ews_pagination_and_half_open_dates():
    client = EWSClient(session=object())
    calls = []

    def call(xml):
        calls.append(xml)
        offset = 0 if len(calls) == 1 else 100
        return ET.fromstring(
            f'<m:R xmlns:m="{NS["m"]}" xmlns:t="{NS["t"]}"><m:RootFolder IncludesLastItemInRange="{"false" if offset==0 else "true"}" IndexedPagingOffset="100"><t:Items><t:Message><t:ItemId Id="{offset}"/></t:Message></t:Items></m:RootFolder></m:R>'
        )

    client.call = call
    ids = client.find_ids(
        '<t:DistinguishedFolderId Id="sentitems"/>', *bounds("2026-08"), sent=True
    )
    assert ids == ["0", "100"]
    assert "item:DateTimeSent" in calls[0]
    assert "2026-08-31T21:00:00Z" in calls[0]
    assert "<t:IsLessThan>" in calls[0]


def test_incomplete_ews_pagination_fails():
    client = EWSClient(session=object())
    client.call = lambda _: ET.fromstring(
        f'<m:R xmlns:m="{NS["m"]}"><m:RootFolder IncludesLastItemInRange="false"/></m:R>'
    )
    with pytest.raises(RuntimeError, match="пагинация"):
        client.find_ids("", *bounds("2026-08"))


def test_api_auth_required(tmp_path):
    from fastapi import HTTPException

    def deny():
        raise HTTPException(401)

    app = FastAPI()
    app.include_router(router(deny, MonthlyService(tmp_path)))
    assert TestClient(app).get("/api/monthly/reports").status_code == 401


def test_api_cannot_approve_missing_export(tmp_path):
    service = MonthlyService(tmp_path)
    report = service.store.create("2026-08", status="draft")
    app = FastAPI()
    app.include_router(router(lambda: True, service))
    response = TestClient(app).post(
        "/api/monthly/reports/" + report["id"] + "/approve", json={"revision": 1}
    )
    assert response.status_code == 409


def test_complete_monthly_pipeline_and_export(tmp_path, monkeypatch):
    from docx import Document
    from src.reporting import pipeline

    # A minimal source template exercises real SQLite, Word export and attachment linkage.
    template = tmp_path / "reference.docx"
    doc = Document()
    doc.sections[0].page_width, doc.sections[0].page_height = 10692000, 7560000
    table = doc.add_table(rows=4, cols=5)
    for cell, text in zip(
        table.rows[0].cells,
        [
            "№",
            "Наименование выполненных мероприятий. Результат работ.",
            "Даты",
            "Ответственный",
            "Договор",
        ],
    ):
        cell.text = text
    for cell, text in zip(
        table.rows[3].cells,
        ["4.2", "Ожидается подтверждение", "Июль", "Солощенко С.С.", "7.2.6"],
    ):
        cell.text = text
    risk = doc.add_table(rows=4, cols=5)
    risk.rows[0].cells[0].text = "Риск-лист"
    risk.rows[1].cells[3].text = "Динамика изменений"
    risk.rows[3].cells[0].text = "2.11"
    doc.save(template)
    service = MonthlyService(tmp_path / "data")
    reference = pipeline.read_reference(template)
    baseline = service.store.create(
        "2026-07",
        status="draft",
        tasks=[task()],
        risks=[],
        template_path=str(template),
        reference=reference,
    )
    baseline = service.export(baseline)
    service.store.update(baseline["id"], status="approved")
    file = tmp_path / "approval.pdf"
    file.write_bytes(b"fixture document")
    message = {
        "id": "m1",
        "date": "2026-08-12T12:00:00Z",
        "subject": "Approval",
        "body": "Планировка согласована",
        "conversation_id": "c1",
        "attachments": [
            {
                "id": "a1",
                "name": "approval.pdf",
                "path": str(file),
                "text": "Планировка согласована",
            }
        ],
    }

    class Mail:
        coverage = {
            "complete": True,
            "folders": [{"name": "Inbox", "found": 1, "read": 1}],
        }

        def collect(self, period, progress):
            return [message]

    class Worker:
        usage = {"runs": 3, "cache_hits": 0}

        def __init__(self, *args):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *args):
            pass

        def ask(self, stage, instruction, payload, schema=None):
            if stage == "triage":
                return {"decisions": {"m1": {"include": True, "reason": "Approval"}}}
            if stage == "extract":
                return {
                    "facts": [
                        {
                            "section": "4.2",
                            "text": "Планировка согласована",
                            "evidence_ids": ["m1"],
                            "attachment_ids": ["a1"],
                        }
                    ]
                }
            if stage == "compose":
                return findings(
                    [
                        dict(
                            task(),
                            id="duplicate-generated-id",
                            previous_id="t1",
                            evidence_ids=["m1"],
                            attachment_ids=["a1"],
                            status="Завершено",
                            result="Планировка согласована оператором.",
                        )
                    ]
                )
            return findings()

    monkeypatch.setattr(pipeline, "EWSClient", Mail)
    monkeypatch.setattr(pipeline, "SDKWorker", Worker)
    monkeypatch.setattr(
        service,
        "prepare_attachments",
        lambda msgs, *args: [service.store.save_message(m) for m in msgs],
    )
    report = service.store.create(
        "2026-08", previous_report_id=baseline["id"], source_type="exchange"
    )
    service.run(report["id"])
    result = service.store.get(report["id"])
    assert result["status"] == "draft", result.get("error")
    assert result["tasks"][0]["id"] == "t1"
    assert result["tasks"][0]["status"] == "Завершено"
    output = Document(result["report_path"])
    assert len(output.tables) == 2
    assert "Планировка согласована оператором" in output.tables[0].rows[2].cells[1].text
    assert list(Path(result["attachments_path"]).rglob("*.pdf"))
    client = TestClient(FastAPI())
    app = FastAPI()
    app.include_router(router(lambda: True, service))
    client = TestClient(app)
    assert (
        client.post(
            f'/api/monthly/reports/{result["id"]}/approve',
            json={"revision": result["revision"], "warnings_acknowledged": True},
        ).status_code
        == 200
    )
    assert (
        client.get(f'/api/monthly/reports/{result["id"]}/document').status_code == 200
    )
    assert (
        client.get(
            f'/api/monthly/reports/{result["id"]}/sources/not-part-of-report'
        ).status_code
        == 404
    )
    assert (
        client.get(f'/api/monthly/reports/{result["id"]}/sources/m1')
        .json()["attachments"][0]
        .get("path")
        is None
    )


def test_sessions_survive_restart_without_plaintext_tokens(tmp_path, monkeypatch):
    from src.reporting.sessions import Sessions

    monkeypatch.setenv("REPORTMASTER_DATA", str(tmp_path))
    first = Sessions()
    first.add("secret-session-token")
    second = Sessions()
    assert "secret-session-token" in second
    assert b"secret-session-token" not in second.path.read_bytes()
    second.discard("secret-session-token")
    assert "secret-session-token" not in first


def test_ews_soap_fault_http500_keeps_exchange_code():
    from types import SimpleNamespace

    response = SimpleNamespace(
        status_code=500,
        content=(
            f'<s:Envelope xmlns:s="{NS["s"]}" xmlns:e="http://schemas.microsoft.com/exchange/services/2006/errors">'
            "<s:Body><s:Fault><faultstring>The request is invalid.</faultstring>"
            "<detail><e:ResponseCode>ErrorInvalidRequest</e:ResponseCode></detail>"
            "</s:Fault></s:Body></s:Envelope>"
        ).encode(),
    )
    client = EWSClient(session=SimpleNamespace(post=lambda *a, **k: response))
    with pytest.raises(RuntimeError, match="ErrorInvalidRequest"):
        client.call("<m:GetItem/>")


def test_ews_reply_reference_uses_item_property():
    client = EWSClient(session=object())

    def call(xml):
        assert 'FieldURI="item:InReplyTo"' in xml
        assert 'FieldURI="message:InReplyTo"' not in xml
        return ET.fromstring(
            f'<m:R xmlns:m="{NS["m"]}" xmlns:t="{NS["t"]}"><m:Items><t:Message>'
            '<t:ItemId Id="item1"/><t:InReplyTo>parent-message</t:InReplyTo>'
            "</t:Message></m:Items></m:R>"
        )

    client.call = call
    assert (
        client.get_items(["item1"], "inbox", False)[0]["in_reply_to"]
        == "parent-message"
    )


def test_model_source_aliases_round_trip():
    from src.reporting.sdk import alias_sources, map_references

    source_id = "a" * 64
    payload = {
        "messages": [{"id": source_id, "body": "Unchanged text"}],
        "evidence_ids": [source_id],
    }
    aliases = alias_sources(payload)
    short = map_references(payload, aliases)
    assert short["messages"][0]["id"] == "ref1"
    assert short["evidence_ids"] == ["ref1"]
    assert map_references(short, {v: k for k, v in aliases.items()}) == payload
    schema = {"required": [source_id], "properties": {source_id: {"type": "boolean"}}}
    assert map_references(schema, aliases)["required"] == ["ref1"]
    assert "ref1" in map_references(schema, aliases)["properties"]


def test_live_worker_permission_error_does_not_mark_failed(tmp_path, monkeypatch):
    path = tmp_path / "history.sqlite3"
    store = ReportStore(path)
    report = store.create("2026-08", worker_pid=12345)

    def inaccessible(*args):
        raise PermissionError("different security context")

    monkeypatch.setattr("src.reporting.store.os.kill", inaccessible)
    assert ReportStore(path).get(report["id"])["status"] == "queued"
