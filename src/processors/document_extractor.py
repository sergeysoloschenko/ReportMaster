"""
Document text extraction for uploaded files and email attachments.
"""

from __future__ import annotations

import hashlib
import logging
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Optional


SUPPORTED_DOCUMENT_EXTENSIONS = {
    ".txt",
    ".csv",
    ".md",
    ".docx",
    ".pdf",
    ".xlsx",
    ".xlsm",
}


@dataclass
class ExtractedDocument:
    filename: str
    extension: str
    content_hash: str
    text: str = ""
    skipped_reason: Optional[str] = None

    @property
    def has_text(self) -> bool:
        return bool(self.text.strip())


class DocumentExtractor:
    """Extract bounded plain text from supported document formats."""

    def __init__(self, max_chars: int = 12000):
        self.logger = logging.getLogger(__name__)
        self.max_chars = max_chars

    def extract_path(self, path: Path) -> ExtractedDocument:
        data = path.read_bytes()
        return self.extract_bytes(path.name, data)

    def extract_bytes(self, filename: str, data: bytes | None) -> ExtractedDocument:
        raw = data or b""
        extension = Path(filename).suffix.lower()
        content_hash = self.hash_bytes(raw)

        if not raw:
            return ExtractedDocument(filename, extension, content_hash, skipped_reason="empty")

        if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
            return ExtractedDocument(filename, extension, content_hash, skipped_reason="unsupported")

        try:
            if extension in {".txt", ".csv", ".md"}:
                text = self._extract_text(raw)
            elif extension == ".docx":
                text = self._extract_docx(raw)
            elif extension == ".pdf":
                text = self._extract_pdf(raw)
            elif extension in {".xlsx", ".xlsm"}:
                text = self._extract_xlsx(raw)
            else:
                text = ""
        except Exception as exc:
            self.logger.warning("Could not extract text from %s: %s", filename, exc)
            return ExtractedDocument(filename, extension, content_hash, skipped_reason=str(exc))

        text = self._normalize_text(text)
        if not text:
            return ExtractedDocument(filename, extension, content_hash, skipped_reason="no_text")

        return ExtractedDocument(filename, extension, content_hash, text=text[: self.max_chars])

    @staticmethod
    def hash_bytes(data: bytes) -> str:
        return hashlib.sha256(data).hexdigest()

    def _extract_text(self, data: bytes) -> str:
        for encoding in ("utf-8", "utf-16", "cp1251", "latin-1"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", errors="ignore")

    def _extract_docx(self, data: bytes) -> str:
        from docx import Document

        doc = Document(BytesIO(data))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        table_cells = []
        for table in doc.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                if values:
                    table_cells.append(" | ".join(values))
        return "\n".join([*paragraphs, *table_cells])

    def _extract_pdf(self, data: bytes) -> str:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
            if sum(len(p) for p in pages) >= self.max_chars:
                break
        return "\n".join(pages)

    def _extract_xlsx(self, data: bytes) -> str:
        import pandas as pd

        sheets = pd.read_excel(BytesIO(data), sheet_name=None, dtype=str)
        chunks = []
        for sheet_name, frame in sheets.items():
            frame = frame.fillna("")
            chunks.append(f"[{sheet_name}]")
            for row in frame.itertuples(index=False, name=None):
                values = [str(value).strip() for value in row if str(value).strip()]
                if values:
                    chunks.append(" | ".join(values))
                if sum(len(chunk) for chunk in chunks) >= self.max_chars:
                    return "\n".join(chunks)
        return "\n".join(chunks)

    @staticmethod
    def _normalize_text(text: str) -> str:
        lines = [" ".join(line.split()) for line in (text or "").splitlines()]
        return "\n".join(line for line in lines if line)
